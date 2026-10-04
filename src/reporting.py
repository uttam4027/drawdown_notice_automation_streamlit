"""Read-side reporting helpers: turn the pipeline's outputs into tables for the UI / audit trail.

Nothing here re-implements business rules. Numbers come from excel_reader + calculation_engine, checks come
from validator.validate_pdf / validate_email, and the QA report is the one written by main.generate_report.
"""
import email
import hashlib
import logging
from email import policy
from pathlib import Path

import pandas as pd

import calculation_engine
import config
import excel_reader
import notice_generator
import template_mapper
import validator
from openpyxl import load_workbook
from utils import email_basename, notice_basename

log = logging.getLogger("drawdown")

QA_REPORT = lambda: config.REPORT_DIR / "final_qa_report.xlsx"      # noqa: E731
ERROR_REPORT = lambda: config.REPORT_DIR / "error_report.xlsx"      # noqa: E731
SUMMARY_FILE = lambda: config.REPORT_DIR / "summary.txt"            # noqa: E731


# ------------------------------------------------------------------ results of the last run
def load_summary():
    """summary.txt -> dict (empty if no run yet). Numeric-looking values stay strings except plain ints."""
    f = SUMMARY_FILE()
    if not f.exists():
        return {}
    out = {}
    for line in f.read_text(encoding="utf-8").splitlines():
        if line.startswith("-") or ":" not in line or line.startswith("BATCH CHECK"):
            continue
        k, v = line.split(":", 1)
        v = v.strip()
        out[k.strip()] = int(v) if v.isdigit() else v
    return out


def load_qa_report():
    f = QA_REPORT()
    return pd.read_excel(f).fillna("") if f.exists() else pd.DataFrame()


def load_error_report():
    f = ERROR_REPORT()
    return pd.read_excel(f).fillna("") if f.exists() else pd.DataFrame()


def output_counts():
    return {"Excel notices": len(list(config.EXCEL_DIR.glob("*.xlsx"))),
            "PDF notices": len(list(config.PDF_DIR.glob("*.pdf"))),
            "Dummy emails": len(list(config.EMAIL_DIR.glob("*.eml"))),
            "QA reports": len(list(config.REPORT_DIR.glob("*.xlsx")))}


def reset_outputs():
    """Remove generated files (excel / pdf / emails / reports) so the demo can be run from a clean state."""
    for folder, pattern in ((config.EXCEL_DIR, "*.xlsx"), (config.PDF_DIR, "*.pdf"),
                            (config.EMAIL_DIR, "*.eml"), (config.REPORT_DIR, "*")):
        folder.mkdir(parents=True, exist_ok=True)
        for f in folder.glob(pattern):
            if f.is_file():
                f.unlink()


# ------------------------------------------------------------------ investor table (from the input workbook)
def load_records(workbook=None):
    df = excel_reader.read_working_data(workbook)
    colmap = excel_reader.build_column_map(df)
    return [excel_reader.get_investor_record(row, colmap) for _, row in df.iterrows()]


def build_investor_table(workbook=None):
    """One row per investor with the figures printed on the notice. Returns (DataFrame, records)."""
    records = load_records(workbook)
    rows = []
    for rec in records:
        calc = calculation_engine.calculate_values(rec)
        try:
            template = template_mapper.get_template(rec["notice_format"], list(config.TEMPLATE_MAPPING.values()))
        except ValueError:
            template = rec["notice_format"] or ""
        rows.append({
            "Folio": rec["folio_no"], "Class": rec["class"],
            "Investor Name": notice_generator.display_name(rec) or "",
            "Notice Format": rec["notice_format"], "Template": template,
            "Drawdown Date": rec["drawdown_date"],
            "Capital Commitment": calc["capital_commitment"],
            "Previous Contribution": calc["capital_contributed"],
            "Drawdown Amount": calc["drawdown_called"],
            "Pending Contribution": calc["pending_previous"],
            "Excess Contribution": calc["excess_contribution"],
            "Amount Payable": calc["amount_payable"],
            "% of Commitment": calc["percentage"],
        })
    return pd.DataFrame(rows), records


def preflight_check(workbook=None):
    """Validate the Excel before running: required columns, template mapping and calculation sanity per row.
    Returns {"investors": n, "issues": [..], "formats": {format: count}}."""
    issues = []
    try:
        records = load_records(workbook)
    except Exception as exc:                                        # noqa: BLE001
        return {"investors": 0, "issues": [f"Cannot read workbook: {exc}"], "formats": {}}
    sheetnames = load_workbook(workbook or config.INPUT_WORKBOOK, read_only=True).sheetnames
    formats = {}
    for rec in records:
        formats[rec["notice_format"]] = formats.get(rec["notice_format"], 0) + 1
        try:
            template_mapper.get_template(rec["notice_format"], sheetnames)
        except ValueError as exc:
            issues.append(f"{rec['folio_no']}: {exc}")
        issues += [f"{rec['folio_no']}: {i}" for i in calculation_engine.validate_financial_calculations(
            rec, calculation_engine.calculate_values(rec))]
    folios = [r["folio_no"] for r in records]
    issues += [f"Duplicate folio {f}" for f in sorted({f for f in folios if folios.count(f) > 1})]
    return {"investors": len(records), "issues": issues, "formats": formats}


# ------------------------------------------------------------------ detailed QA (re-uses validator functions)
def read_eml(path):
    """Parse a generated .eml for display: dict with to, subject, from, body, attachment name/size."""
    with open(path, "rb") as fh:
        msg = email.message_from_binary_file(fh, policy=policy.default)
    atts = list(msg.iter_attachments())
    body = msg.get_body(preferencelist=("plain",))
    return {"from": msg["From"], "to": msg["To"], "subject": msg["Subject"], "date": msg["Date"],
            "body": body.get_content() if body else "",
            "attachment": atts[0].get_filename() if atts else "",
            "attachment_kb": round(len(atts[0].get_content()) / 1024, 1) if atts else 0,
            "demo_header": msg["X-Demo-Mode"]}


def paths_for(rec, template):
    stem = notice_basename(rec["folio_no"], template)
    return (config.PDF_DIR / f"{stem}.pdf",
            config.EMAIL_DIR / f"{email_basename(rec['folio_no'], template)}.eml")


def build_qa_detail(workbook=None):
    """Re-run validator.validate_pdf / validate_email over the generated files to get check-level results.
    Returns (per-investor DataFrame, check summary DataFrame)."""
    qa = load_qa_report()
    if qa.empty:
        return pd.DataFrame(), pd.DataFrame()
    records = {r["folio_no"]: r for r in load_records(workbook)}
    sheetnames = list(config.TEMPLATE_MAPPING.values())
    hashes, rows = {}, []
    for _, q in qa.iterrows():
        rec = records.get(q["Folio No."])
        row = {"Folio": q["Folio No."], "Template": q["Format"], "PDF Status": "FAIL", "Email Status": "FAIL",
               "Calculation Status": q["Financial Validation"] or "FAIL",
               "QA Status": "PASS" if q["Status"] == "SUCCESS" else "FAIL", "Detail": q["Error Message"], "_checks": {}}
        if rec and q["PDF File"]:
            template = template_mapper.get_template(rec["notice_format"], sheetnames)
            calc = calculation_engine.calculate_values(rec)
            pdf, eml = paths_for(rec, template)
            if pdf.exists():
                res = validator.validate_pdf(pdf, rec, calc, int(q["Expected Pages"]), pdf.stem)
                row["_checks"] = dict(res["checks"], missing_pdf="PASS")
                row["PDF Status"] = "PASS" if not res["problems"] else "FAIL"
                hashes.setdefault(hashlib.md5(pdf.read_bytes()).hexdigest(), []).append(pdf.name)
            else:
                row["_checks"]["missing_pdf"] = "FAIL"
            if eml.exists() and pdf.exists():
                probs = validator.validate_email(eml, pdf, rec, config_to(rec), config.EMAIL_SUBJECT.format(
                    folio=rec["folio_no"]))
                row["Email Status"] = "PASS" if not probs else "FAIL"
                row["_checks"]["attachment"] = "FAIL" if any("attach" in p for p in probs) else "PASS"
        rows.append(row)
    dup = {n for names in hashes.values() if len(names) > 1 for n in names}
    for row in rows:
        row["_checks"]["duplicate_pdf"] = "FAIL" if any(n.startswith(row["Folio"] + "_") for n in dup) else "PASS"
        row["_checks"].setdefault("missing_pdf", "FAIL")
    for row in rows:       # the report is a snapshot of the run; files changed since then must also count
        if "FAIL" in (row["PDF Status"], row["Email Status"], row["Calculation Status"]):
            row["QA Status"] = "FAIL"
    detail = pd.DataFrame(rows)

    labels = [  # (label shown, validator check key, what it catches)
        ("#REF!", "error_values", "Excel error text in the PDF (one scan covers #REF!, #VALUE!, #NAME?, #DIV/0!)"),
        ("#VALUE!", "error_values", "same scan as above"),
        ("#NAME?", "error_values", "same scan as above"),
        ("#DIV/0!", "error_values", "same scan as above"),
        ("Cross-sheet references", "cross_sheet_references", "visible 'Working!' / 'Format - ...' references"),
        ('Blank "To" section', "to_section", 'the "To," block has investor name and address'),
        ("Placeholder text", "banned_text", "abc@gmail.com, CLIENT NAME AND ADDRESS, build notes"),
        ("Signature placeholder", "signature", "For <client> / On behalf of / Sd/- / Designated Partner present"),
        ("Missing PDF", "missing_pdf", "a PDF exists for the investor"),
        ("Wrong attachment", "attachment", "email carries exactly this investor's PDF, byte-identical"),
        ("Duplicate PDF", "duplicate_pdf", "no two investors share identical PDF content"),
    ]
    summ = []
    for label, key, what in labels:
        vals = [r["_checks"].get(key) for _, r in detail.iterrows() if r["_checks"].get(key)]
        failed = sum(v == "FAIL" for v in vals)
        summ.append({"Check": label, "What it catches": what, "Checked": len(vals), "Failed": failed,
                     "Result": "PASS" if vals and not failed else ("FAIL" if failed else "n/a")})
    return detail.drop(columns=["_checks"]), pd.DataFrame(summ)


def config_to(rec):
    from email_service import recipient_address
    return recipient_address(rec)
