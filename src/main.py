"""Orchestration of the full drawdown-notice workflow (STEP 1-16).

Every investor is processed inside its own try/except: one failure never stops the run, it is written to
error_report.xlsx and the next investor is processed.
"""
import logging
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import config
import excel_reader
import template_mapper
import calculation_engine
import notice_generator
import pdf_generator
import email_service
import validator
from utils import email_basename, now_stamp, notice_basename, setup_logging

QA_COLUMNS = ["Folio No.", "Class", "Format", "Investor Name", "Email", "PDF File", "Email File", "Pages",
              "Broken References", "Financial Validation", "Placeholder Check", "To Section", "Client Email",
              "Signature", "Status", "Error Message", "Amount Payable", "Expected Pages", "Data Warnings",
              "QA Warnings"]
ERROR_COLUMNS = ["Folio No.", "Class", "Format", "Error", "Stage", "Timestamp"]


def _clean_outputs():
    """Remove files from a previous run so stale PDFs/emails can never be mixed with new ones."""
    for folder, pattern in ((config.EXCEL_DIR, "*.xlsx"), (config.PDF_DIR, "*.pdf"), (config.EMAIL_DIR, "*.eml")):
        folder.mkdir(parents=True, exist_ok=True)
        for f in folder.glob(pattern):
            f.unlink()
    config.REPORT_DIR.mkdir(parents=True, exist_ok=True)


# Pipeline steps shown by a UI, and the error-row stage that maps to each step
STAGES = ["Excel Input", "Data Validation", "Template Mapping", "Financial Calculation", "Notice Generation",
          "PDF Generation", "PDF QA", "Dummy Email", "Audit Report"]
_ERROR_STAGE_TO_STEP = {"template selection": "Template Mapping", "calculation": "Financial Calculation",
                        "notice generation": "Notice Generation", "pdf export": "PDF Generation",
                        "pdf qa": "PDF QA", "qa": "PDF QA", "email creation": "Dummy Email"}


def _emit(progress, stage, status, done=0, total=0, folio=""):
    """Optional progress hook for a UI (Streamlit). status: running | done | failed. Never raises."""
    if progress is None:
        return
    try:
        progress(stage, status, done, total, folio)
    except Exception:                                               # noqa: BLE001 - UI must not break the run
        logging.getLogger("drawdown").exception("progress callback failed")


def _step_status(error_rows, step):
    return "failed" if any(_ERROR_STAGE_TO_STEP.get(e["Stage"]) == step for e in error_rows) else "done"


def _passfail(ok):
    return "PASS" if ok else "FAIL"


def generate_report(qa_rows, error_rows, started, batch_facts=None, batch_fails=None):
    """STEP 15: final_qa_report.xlsx, error_report.xlsx, summary.txt. Returns the summary dict."""
    qa = pd.DataFrame(qa_rows, columns=QA_COLUMNS)
    err = pd.DataFrame(error_rows, columns=ERROR_COLUMNS)
    _write_xlsx(qa, config.REPORT_DIR / "final_qa_report.xlsx", "QA")
    _write_xlsx(err, config.REPORT_DIR / "error_report.xlsx", "Errors")

    total = len(qa)
    ok = int((qa["Status"] == "SUCCESS").sum())
    summary = {
        "Total investors": total,
        "Processed": total,
        "Successful": ok,
        "Failed": total - ok,
        "PDFs": int(qa["PDF File"].astype(bool).sum()),
        "Dummy emails": int(qa["Email File"].astype(bool).sum()),
        "Excel notices": (batch_facts or {}).get("Excel notices", ""),
        "Templates processed": (batch_facts or {}).get("Templates processed", ""),
        "QA failures": int((qa["Status"] == "QA_FAIL").sum()),
        "Errors (exceptions)": int((qa["Status"] == "ERROR").sum()),
        "Investors with blank name in Working": int(qa["Data Warnings"].str.contains("name blank").sum()),
        "Investors with blank email in Working": int(qa["Data Warnings"].str.contains("email blank").sum()),
        "PDFs with QA warnings": int((qa["QA Warnings"].astype(str).str.len() > 0).sum()),
        "Batch verification": "PASS" if not batch_fails else "FAIL",
        "Run time (s)": round(time.time() - started, 1),
    }
    lines = [f"{k}: {v}" for k, v in summary.items()]
    for tpl, pg in (batch_facts or {}).get("Pages per template", {}).items():
        lines.append(f"Pages - {tpl}: {pg}")
    lines += [f"BATCH CHECK FAILED: {f}" for f in (batch_fails or [])]
    lines += ["", "Known warnings (data not in workbook):"] + [f"- {w}" for w in config.KNOWN_WARNINGS]
    (config.REPORT_DIR / "summary.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return summary


def _write_xlsx(df, path, sheet):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=sheet, index=False)
        ws = xw.sheets[sheet]
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 60)
        ws.freeze_panes = "A2"


def main(argv=None, progress=None):
    """progress: optional callable(stage, status, done, total, folio) - used by the Streamlit UI."""
    argv = argv or []
    started = time.time()
    log = setup_logging()
    _clean_outputs()

    print("=" * 49)
    print("DRAWDOWN NOTICE AUTOMATION")
    print("=" * 49)

    # STEP 1-4 -------------------------------------------------------------------------------
    _emit(progress, "Excel Input", "running")
    df = excel_reader.read_working_data()
    _emit(progress, "Excel Input", "done", len(df), len(df))
    _emit(progress, "Data Validation", "running")
    colmap = excel_reader.build_column_map(df)                      # STEP 3: validate required columns
    records = [excel_reader.get_investor_record(row, colmap) for _, row in df.iterrows()]
    log.info("Number of investors: %d", len(records))
    sheetnames = load_workbook(config.INPUT_WORKBOOK, read_only=True).sheetnames

    # A1-001 test before processing everybody ---------------------------------------------------
    if "--skip-test" not in argv:
        fails = validator.run_a1_001_calc_test(records)
        if fails:
            print("A1-001 calculation test FAILED:\n  " + "\n  ".join(fails))
            log.error("A1-001 calculation test failed: %s", fails)
            _emit(progress, "Data Validation", "failed")
            return 1
        print("A1-001 calculation test: PASS")
    _emit(progress, "Data Validation", "done", len(records), len(records))

    # STEP 5-10: build one self-contained notice workbook per investor ------------------------------
    built, qa_rows, error_rows = [], {}, []

    def fail(rec, stage, exc):
        log.error("%s ERROR at %s: %s", rec["folio_no"], stage, exc)
        error_rows.append({"Folio No.": rec["folio_no"], "Class": rec["class"], "Format": rec["notice_format"],
                           "Error": str(exc), "Stage": stage, "Timestamp": now_stamp()})
        qa_rows[rec["folio_no"]] = {
            "Folio No.": rec["folio_no"], "Class": rec["class"], "Format": rec["notice_format"],
            "Investor Name": notice_generator.display_name(rec) or "", "Email": rec["email"], "PDF File": "", "Email File": "",
            "Pages": "", "Broken References": "", "Financial Validation": "", "Placeholder Check": "",
            "To Section": "", "Client Email": "", "Signature": "", "QA Warnings": "",
            "Status": "ERROR", "Error Message": f"[{stage}] {exc}", "Amount Payable": "",
            "Expected Pages": "", "Data Warnings": "; ".join(notice_generator.data_warnings(rec))}

    for step in ("Template Mapping", "Financial Calculation"):
        _emit(progress, step, "running")
    for i, rec in enumerate(records, 1):
        _emit(progress, "Notice Generation", "running", i, len(records), rec["folio_no"])
        stage = "template selection"
        try:
            sheet = template_mapper.get_template(rec["notice_format"], sheetnames)
            stage = "calculation"
            calc = calculation_engine.calculate_values(rec)
            fin_issues = calculation_engine.validate_financial_calculations(rec, calc)
            stage = "notice generation"
            stem = notice_basename(rec["folio_no"], sheet)
            info = notice_generator.generate_notice(rec, calc, sheet, config.EXCEL_DIR / f"{stem}.xlsx")
            built.append({"rec": rec, "calc": calc, "sheet": sheet, "stem": stem, "xlsx": info["path"],
                          "expected_pages": info["expected_pages"], "fin_issues": fin_issues})
        except Exception as exc:                                    # noqa: BLE001 - never stop the whole run
            fail(rec, stage, exc)

    # STEP 11: PDFs (batched through one LibreOffice session) ----------------------------------------
    for step in ("Template Mapping", "Financial Calculation", "Notice Generation"):
        _emit(progress, step, _step_status(error_rows, step), len(built), len(records))
    _emit(progress, "PDF Generation", "running", 0, len(built))
    pdfs = pdf_generator.export_pdf([b["xlsx"] for b in built]) if built else {}
    _emit(progress, "PDF Generation", "failed" if any(v is None for v in pdfs.values()) else "done",
          len(built), len(built))

    # STEP 12-14: QA, dummy email, attach PDF ----------------------------------------------------------
    _emit(progress, "PDF QA", "running")
    _emit(progress, "Dummy Email", "running")
    for i, b in enumerate(built, 1):
        rec, calc, sheet, stem = b["rec"], b["calc"], b["sheet"], b["stem"]
        _emit(progress, "PDF QA", "running", i, len(built), rec["folio_no"])
        stage = "pdf export"
        try:
            pdf = pdfs.get(stem)
            if pdf is None:
                raise RuntimeError("LibreOffice did not produce a PDF")
            stage = "pdf qa"
            qa = validator.validate_pdf(pdf, rec, calc, b["expected_pages"], stem)
            stage = "email creation"
            eml = config.EMAIL_DIR / f"{email_basename(rec['folio_no'], sheet)}.eml"
            to_addr = email_service.create_dummy_email(rec, pdf, eml)
            email_problems = validator.validate_email(
                eml, pdf, rec, to_addr, config.EMAIL_SUBJECT.format(folio=rec["folio_no"]))

            problems = qa["problems"] + b["fin_issues"] + email_problems
            status = "SUCCESS" if not problems else "QA_FAIL"
            c = qa["checks"]
            qa_rows[rec["folio_no"]] = {
                "Folio No.": rec["folio_no"], "Class": rec["class"], "Format": rec["notice_format"],
                "Investor Name": notice_generator.display_name(rec) or "", "Email": to_addr, "PDF File": pdf.name,
                "Email File": eml.name, "Pages": qa["pages"], "Broken References": c["broken_references"],
                "Financial Validation": _passfail(not b["fin_issues"] and c["financial_values"] == "PASS"),
                "Placeholder Check": _passfail(c["banned_text"] == "PASS" and c["client_block"] == "PASS"),
                "To Section": c["to_section"], "Client Email": c["client_email"], "Signature": c["signature"],
                "Status": status, "Error Message": " | ".join(problems),
                "Amount Payable": calc["amount_payable"], "Expected Pages": b["expected_pages"],
                "Data Warnings": "; ".join(notice_generator.data_warnings(rec)),
                "QA Warnings": " | ".join(qa["warnings"])}
            if problems:
                error_rows.append({"Folio No.": rec["folio_no"], "Class": rec["class"],
                                   "Format": rec["notice_format"], "Error": " | ".join(problems),
                                   "Stage": "qa", "Timestamp": now_stamp()})
            log.info("%s %s", rec["folio_no"], status)
        except Exception as exc:                                    # noqa: BLE001
            fail(rec, stage, exc)

    for step in ("PDF Generation", "PDF QA", "Dummy Email"):
        _emit(progress, step, _step_status(error_rows, step), len(built), len(built))

    # keep the report in Working-sheet order
    ordered = [qa_rows[r["folio_no"]] for r in records if r["folio_no"] in qa_rows]
    for row in ordered:
        print(f"{row['Folio No.']:<8} {row['Format']:<34} {row['Status']}"
              + (f"  <- {row['Error Message'][:80]}" if row["Status"] != "SUCCESS" else ""))

    # Batch-level verification (counts, one email per investor, all templates used, pages per template) ----------
    batch_fails, batch_facts = validator.verify_batch(records, ordered)
    for f in batch_fails:
        log.error("BATCH CHECK FAILED: %s", f)
        print("BATCH CHECK FAILED:", f)

    # A1-001 output test --------------------------------------------------------------------------
    out_fails = validator.check_a1_001_outputs() if "--skip-test" not in argv else []
    if out_fails:
        log.error("A1-001 output test failed: %s", out_fails)
        print("A1-001 output test FAILED:\n  " + "\n  ".join(out_fails))
    elif "--skip-test" not in argv:
        print("A1-001 output test: PASS")

    # STEP 15-16 ------------------------------------------------------------------------------------
    _emit(progress, "Audit Report", "running")
    summary = generate_report(ordered, error_rows, started, batch_facts, batch_fails)
    _emit(progress, "Audit Report", "failed" if (out_fails or batch_fails) else "done")
    print("\n" + "=" * 49)
    print("Total investors:", summary["Total investors"])
    print("Processed:", summary["Processed"], " Successful:", summary["Successful"], " Failed:", summary["Failed"])
    print("PDF generated:", summary["PDFs"], " Dummy emails generated:", summary["Dummy emails"])
    print("QA errors:", summary["QA failures"] + summary["Errors (exceptions)"])
    print("Excel notices:", batch_facts["Excel notices"], " Templates processed:", batch_facts["Templates processed"])
    print("Batch verification:", "PASS" if not batch_fails else f"FAIL ({len(batch_fails)})")
    print("Output folder:", config.OUT_DIR)
    print("=" * 49)
    return 0 if not (out_fails or batch_fails) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
